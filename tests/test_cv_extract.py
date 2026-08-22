import os, sys, tempfile, unittest

HERE = os.path.dirname(__file__)
CORE = os.path.join(HERE, "..", "addon", "globalPlugins", "jobFormFiller", "core")
sys.path.insert(0, os.path.abspath(CORE))

import cvparse
from matcher import FieldDescriptor as F, match_field

CV_EN = [
    "Mohammed Al Omrani",
    "Bristol, UK",
    "example@example.com  +44 7700 900123",
    "Education",
    "BA (Hons) Education, University of the West of England, 2026",
    "Experience",
    "Technology Support Volunteer, Sight Support West of England",
]


def make_docx(path):
    import docx
    d = docx.Document()
    for line in CV_EN:
        d.add_paragraph(line)
    d.save(path)


def make_pdf(path):
    from fpdf import FPDF, XPos, YPos
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in CV_EN:
        pdf.cell(0, 8, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(path)


class TestDocxPdfExtraction(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def test_docx_roundtrip(self):
        p = os.path.join(self.dir, "cv.docx")
        make_docx(p)
        text = cvparse.extract_text(p)
        r = cvparse.parse_cv_text(text)
        self.assertEqual(r["email"], "example@example.com")
        self.assertEqual(r["full_name"], "Mohammed Al Omrani")
        self.assertIn("education", r)
        self.assertIn("experience", r)

    def test_pdf_roundtrip(self):
        p = os.path.join(self.dir, "cv.pdf")
        make_pdf(p)
        text = cvparse.extract_text(p)
        self.assertIn("example@example.com", text)
        r = cvparse.parse_cv_text(text)
        self.assertEqual(r["email"], "example@example.com")
        self.assertIn("education", r)

    def test_unsupported_format_raises(self):
        with self.assertRaises(NotImplementedError):
            cvparse.extract_text("cv.rtf")


class TestMultilingualParsing(unittest.TestCase):
    FR = (
        "Jean Dupont\n"
        "Courriel : jean.dupont@example.fr\n"
        "Formation\n"
        "Master en informatique, Universite de Paris, 2025\n"
        "Experience professionnelle\n"
        "Developpeur, ACME SARL\n"
    )

    def test_french_sections(self):
        r = cvparse.parse_cv_text(self.FR)
        self.assertEqual(r["email"], "jean.dupont@example.fr")
        self.assertIn("education", r)      # "Formation"
        self.assertIn("experience", r)     # "Experience professionnelle"

    def test_accented_labels_match(self):
        # accented, real-world field labels must still resolve
        self.assertEqual(match_field(F(label="Téléphone")).key, "phone")
        self.assertEqual(match_field(F(label="Prénom")).key, "given_name")
        self.assertEqual(match_field(F(label="Dirección")).key, "address_line1")
        self.assertEqual(match_field(F(label="Numéro de téléphone")).key, "phone")


if __name__ == "__main__":
    unittest.main(verbosity=2)
